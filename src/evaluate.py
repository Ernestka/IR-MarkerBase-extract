"""
Compare the pipeline's final table against a hand-extracted GOLD-STANDARD table
(step 4 of the work plan). Both use the final-table columns (export.COLUMNS).

    python src/evaluate.py --template gold.xlsx           # empty gold sheet to fill in
    python src/evaluate.py gold.xlsx                      # vs data/final/ir_extraction.csv
    python src/evaluate.py gold.csv local_output/final/ir_extraction.csv --out eval/

Filling the gold sheet: one row per data point, exactly like the final table —
a genotype row (Gene + Variant/marker filled), a bioassay row (Insecticide name
filled, no Gene) or a geno-pheno row (Gene + RR/RS/SS survivors/dead filled).
`Record type` may be left empty (it is inferred). Leave a cell empty or NA when
the paper does not report it. Papers are matched by DOI (else Paper/Study ID).

What it reports:
  records  recall    = gold rows found in the extraction
           precision = extracted rows that exist in the gold set (only for papers
                       that are in the gold set)
  fields   for every matched row and every column the gold row fills: is the
           extracted value the same? (counts exact; proportions +-0.01;
           coordinates +-0.05 deg; text compared case/space/punctuation-blind)
Writes <out>/eval_fields.csv, eval_records.csv and eval_mismatches.csv.
"""
import argparse
import csv
import difflib
import re
import sys
from pathlib import Path

import export
import targets

ROOT = Path(__file__).resolve().parent.parent

# Bookkeeping columns: not something the extraction can get right or wrong.
SKIP = {"Paper/Study ID", "Record type", "Survey ID", "Human verified", "Coordinates verified",
        "AI extraction confidence", "Source sentence/table", "Data quality flag",
        "Evidence location (data source)", "Evidence location (coordinates)"}
PROPORTIONS = {"Allele frequency", "RR frequency", "RS frequency", "SS frequency",
               "Survival proportion", "Mortality proportion"}
COORDS = {"Latitude", "Longitude"}


# --- Reading ----------------------------------------------------------------

def read_table(path):
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        head = [str(h).strip() if h is not None else "" for h in rows[0]]
        return [{h: ("" if v is None else str(v)) for h, v in zip(head, r)} for r in rows[1:]
                if any(v not in (None, "") for v in r)]
    with path.open(newline="", encoding="utf-8-sig") as f:
        return [{(k or "").strip(): v for k, v in r.items()} for r in csv.DictReader(f)]


def _empty(v):
    return v is None or str(v).strip() in ("", "NA", "N/A", "na", "-")


def _num(v):
    try:
        return float(str(v).replace(",", ".").replace("%", ""))
    except (TypeError, ValueError):
        return None


def _text(v):
    return re.sub(r"[\W_]+", "", str(v).lower())


def record_type(r):
    t = (r.get("Record type") or "").strip()
    if t and not _empty(t):
        return t
    if not _empty(r.get("Gene")) and any(not _empty(r.get(c)) for c in
                                         ("RR survivors", "RS survivors", "SS survivors",
                                          "RR dead", "RS dead", "SS dead", "Phenotype group")):
        return "geno_pheno"
    if not _empty(r.get("Gene")) or not _empty(r.get("Variant/marker")):
        return "genotype"
    if not _empty(r.get("Insecticide name")):
        return "bioassay"
    return "other"


def _paper(r):
    doi = (r.get("DOI") or "").strip().lower().removeprefix("https://doi.org/")
    return doi if not _empty(doi) else _text(r.get("Paper/Study ID") or "")


def _marker(r):
    t = targets.normalise(r.get("Gene") or "", r.get("Variant/marker") or "")
    return f"{t['gene']} {t['variant']}" if t else _text(r.get("Variant/marker") or "")


def _key(r):
    """What identifies a data point within a paper (site + test/marker)."""
    parts = [_text(r.get("Sampling site") or "")]
    rt = record_type(r)
    if rt in ("bioassay", "geno_pheno"):
        parts += [_text(r.get(c) or "") for c in ("Insecticide name", "Insecticide concentration")]
        syn = _text(r.get("Synergist used") or "")
        parts.append("" if syn in ("", "none", "na") else syn)
    if rt in ("genotype", "geno_pheno"):
        parts.append(_marker(r))
    parts.append(_text(r.get("Molecular species") or r.get("Species reported") or ""))
    parts.append((r.get("Collection start date") or "")[:4])
    return parts


# --- Matching ---------------------------------------------------------------

def _similarity(a, b):
    """0-1: how alike two keys are. The test/marker parts must match exactly."""
    if a[1:-2] != b[1:-2]:
        return 0.0
    site = difflib.SequenceMatcher(None, a[0], b[0]).ratio() if a[0] or b[0] else 1.0
    extra = sum(x == y for x, y in zip(a[-2:], b[-2:])) / 2
    return 0.8 * site + 0.2 * extra


def match(gold, ext, threshold=0.6):
    """Greedy one-to-one matching of gold rows to extracted rows within each
    (paper, record type). -> [(gold_row, ext_row or None)], unmatched ext rows."""
    pairs, used = [], set()
    by_group = {}
    for i, e in enumerate(ext):
        by_group.setdefault((_paper(e), record_type(e)), []).append(i)
    cands = []
    for gi, g in enumerate(gold):
        for ei in by_group.get((_paper(g), record_type(g)), []):
            s = _similarity(_key(g), _key(ext[ei]))
            if s >= threshold:
                cands.append((s, gi, ei))
    best = {}
    for s, gi, ei in sorted(cands, reverse=True):
        if gi not in best and ei not in used:
            best[gi] = ei
            used.add(ei)
    for gi, g in enumerate(gold):
        pairs.append((g, ext[best[gi]] if gi in best else None))
    gold_papers = {_paper(g) for g in gold}
    extra = [e for i, e in enumerate(ext) if i not in used and _paper(e) in gold_papers]
    return pairs, extra


def same(col, gv, ev):
    if _empty(ev):
        return False
    g, e = _num(gv), _num(ev)
    if g is not None and e is not None:
        if col in COORDS:
            return abs(g - e) <= 0.05
        if col in PROPORTIONS:
            if g > 1:
                g = g / 100                      # gold typed as a percentage
            return abs(g - e) <= 0.01
        return abs(g - e) < 1e-9
    if col == "Gene":
        return _text(targets.canonical_gene(gv)) == _text(targets.canonical_gene(ev))
    if col == "Variant/marker":
        return _marker({"Gene": "", "Variant/marker": gv}) == _marker({"Gene": "", "Variant/marker": ev})
    return _text(gv) == _text(ev)


# --- Report -----------------------------------------------------------------

def evaluate(gold_path, ext_path, out_dir):
    gold, ext = read_table(gold_path), read_table(ext_path)
    pairs, extra = match(gold, ext)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    rec = {}
    for g, e in pairs:
        t = rec.setdefault(record_type(g), {"gold": 0, "found": 0, "extra": 0})
        t["gold"] += 1
        t["found"] += e is not None
    for e in extra:
        rec.setdefault(record_type(e), {"gold": 0, "found": 0, "extra": 0})["extra"] += 1

    fields, mismatches = {}, []
    for g, e in pairs:
        if e is None:
            mismatches.append({"paper": _paper(g), "record type": record_type(g),
                               "key": " | ".join(_key(g)), "column": "(whole row)",
                               "gold": "present", "extracted": "MISSING"})
            continue
        for col in export.COLUMNS:
            if col in SKIP or _empty(g.get(col)):
                continue
            ok = same(col, g[col], e.get(col))
            f = fields.setdefault(col, [0, 0])
            f[0] += 1
            f[1] += ok
            if not ok:
                mismatches.append({"paper": _paper(g), "record type": record_type(g),
                                   "key": " | ".join(_key(g)), "column": col,
                                   "gold": g[col], "extracted": e.get(col, "")})
    for e in extra:
        mismatches.append({"paper": _paper(e), "record type": record_type(e),
                           "key": " | ".join(_key(e)), "column": "(whole row)",
                           "gold": "absent", "extracted": "EXTRA ROW"})

    def pct(a, b):
        return f"{100 * a / b:.0f}%" if b else "—"

    with (out_dir / "eval_records.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["record type", "gold rows", "found (recall)", "extra rows", "precision"])
        print(f"\n{'record type':<12} {'gold':>5} {'found':>6} {'recall':>7} {'extra':>6} {'precision':>9}")
        for t, v in sorted(rec.items()):
            found_all = v["found"] + v["extra"]
            row = [t, v["gold"], v["found"], v["extra"], pct(v["found"], found_all)]
            w.writerow(row)
            print(f"{t:<12} {v['gold']:>5} {v['found']:>6} {pct(v['found'], v['gold']):>7} "
                  f"{v['extra']:>6} {pct(v['found'], found_all):>9}")

    with (out_dir / "eval_fields.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["column", "compared", "correct", "accuracy"])
        print(f"\n{'column':<38} {'compared':>8} {'correct':>8} {'accuracy':>8}")
        for col in export.COLUMNS:
            if col in fields:
                n, ok = fields[col]
                w.writerow([col, n, ok, pct(ok, n)])
                print(f"{col:<38} {n:>8} {ok:>8} {pct(ok, n):>8}")

    with (out_dir / "eval_mismatches.csv").open("w", newline="", encoding="utf-8") as f:
        cols = ["paper", "record type", "key", "column", "gold", "extracted"]
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(mismatches)
    print(f"\n{len(mismatches)} mismatch(es) -> {out_dir / 'eval_mismatches.csv'}")


def write_template(path):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "gold"
        ws.append(export.COLUMNS)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "B2"
        wb.save(path)
    else:
        with path.open("w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerow(export.COLUMNS)
    print(f"Wrote empty gold-standard sheet: {path}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("gold", nargs="?", help="hand-extracted gold table (.csv or .xlsx)")
    ap.add_argument("extracted", nargs="?", default=str(ROOT / "data" / "final" / "ir_extraction.csv"))
    ap.add_argument("--out", default=str(ROOT / "local_output" / "eval"))
    ap.add_argument("--template", metavar="PATH", help="write an empty gold sheet and exit")
    args = ap.parse_args()
    if args.template:
        return write_template(args.template)
    if not args.gold:
        sys.exit("Give the gold table, or --template PATH to create one.")
    evaluate(args.gold, args.extracted, args.out)


if __name__ == "__main__":
    main()
